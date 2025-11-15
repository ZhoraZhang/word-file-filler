import os
import re
import traceback
import pandas as pd
from docx import Document
import tkinter as tk
from tkinter import filedialog, messagebox, scrolledtext


# ================= 占位符替换核心逻辑 =================

# 正则：匹配 {{名字}} 或 {名字}
PLACEHOLDER_PATTERN = re.compile(r"\{\{([^{}]+)\}\}|\{([^{}]+)\}")


def replace_placeholders_in_text(text, row):
    """
    在一段纯文本中，用当前 Excel 行 row 替换所有占位符。
    占位符形式：
      - {{列名}}
      - {列名}
    """
    def repl(match):
        # 两种括号形式取其一
        name = match.group(1) or match.group(2)
        name = str(name).strip()

        if name in row.index and pd.notna(row[name]):
            return str(row[name])
        else:
            return ""

    return PLACEHOLDER_PATTERN.sub(repl, text)


def replace_in_paragraph(paragraph, row):
    """
    支持占位符被拆成多个 run：
    1. 把所有 run 拼成完整字符串
    2. 字符串里执行占位符替换
    3. 如果有变化，用一个 run 写回
    """
    if not paragraph.runs:
        return

    full_text = "".join(run.text for run in paragraph.runs)
    new_text = replace_placeholders_in_text(full_text, row)

    if new_text == full_text:
        return  # 没有占位符，无需修改

    paragraph.runs[0].text = new_text
    for run in paragraph.runs[1:]:
        run.text = ""


def replace_in_document(doc, row):
    """
    在整个文档（段落 + 表格里的段落）中进行占位符替换
    """
    # 普通段落
    for p in doc.paragraphs:
        replace_in_paragraph(p, row)

    # 表格中的段落
    for table in doc.tables:
        for tr in table.rows:
            for cell in tr.cells:
                for p in cell.paragraphs:
                    replace_in_paragraph(p, row)


def generate_reports(template_path, excel_path, output_dir, log_func=print):
    """
    核心函数：
    - 读取 excel_path
    - 对每一行，用 template_path 生成一份报告，保存到 output_dir
    - log_func 用于输出日志（可以是 print，也可以是 GUI 的日志窗口）
    """
    if not os.path.isfile(template_path):
        raise FileNotFoundError(f"模板文件不存在: {template_path}")
    if not os.path.isfile(excel_path):
        raise FileNotFoundError(f"Excel 文件不存在: {excel_path}")
    os.makedirs(output_dir, exist_ok=True)

    log_func(f"读取 Excel: {excel_path}")
    df = pd.read_excel(excel_path)

    if df.empty:
        log_func("警告：Excel 中没有数据。")
        return

    log_func(f"共读取到 {len(df)} 行数据。")

    for idx, row in df.iterrows():
        doc = Document(template_path)

        replace_in_document(doc, row)

        # 输出文件名：优先尝试几个常见列名
        name = None
        for key in ["姓名", "拜访人姓名", "拜访人", "Name"]:
            if key in row.index and pd.notna(row[key]):
                name = str(row[key])
                break
        if name is None:
            name = f"report_{idx + 1}"

        out_path = os.path.join(output_dir, f"{name}.docx")
        doc.save(out_path)
        log_func(f"生成报告: {out_path}")

    log_func("全部生成完成。")


# ================= GUI 部分 =================

class ReportFillerApp:
    def __init__(self, master):
        self.master = master
        master.title("Word 报告自动填充工具")
        master.geometry("700x450")

        # 模板路径
        self.label_template = tk.Label(master, text="模板文件（.docx）：")
        self.label_template.grid(row=0, column=0, padx=10, pady=10, sticky="e")

        self.entry_template = tk.Entry(master, width=60)
        self.entry_template.grid(row=0, column=1, padx=5, pady=10, sticky="w")

        self.btn_template = tk.Button(master, text="浏览...", command=self.browse_template)
        self.btn_template.grid(row=0, column=2, padx=10, pady=10)

        # Excel 路径
        self.label_excel = tk.Label(master, text="Excel 文件（.xlsx/.xls）：")
        self.label_excel.grid(row=1, column=0, padx=10, pady=10, sticky="e")

        self.entry_excel = tk.Entry(master, width=60)
        self.entry_excel.grid(row=1, column=1, padx=5, pady=10, sticky="w")

        self.btn_excel = tk.Button(master, text="浏览...", command=self.browse_excel)
        self.btn_excel.grid(row=1, column=2, padx=10, pady=10)

        # 输出文件夹
        self.label_output = tk.Label(master, text="输出文件夹：")
        self.label_output.grid(row=2, column=0, padx=10, pady=10, sticky="e")

        self.entry_output = tk.Entry(master, width=60)
        self.entry_output.grid(row=2, column=1, padx=5, pady=10, sticky="w")

        self.btn_output = tk.Button(master, text="浏览...", command=self.browse_output)
        self.btn_output.grid(row=2, column=2, padx=10, pady=10)

        # 开始按钮
        self.btn_run = tk.Button(master, text="开始生成", width=15, command=self.run)
        self.btn_run.grid(row=3, column=1, pady=10)

        # 日志区域
        self.log_text = scrolledtext.ScrolledText(master, width=80, height=15)
        self.log_text.grid(row=4, column=0, columnspan=3, padx=10, pady=10)

    def browse_template(self):
        path = filedialog.askopenfilename(
            title="选择 Word 模板文件",
            filetypes=[("Word 文件", "*.docx"), ("所有文件", "*.*")]
        )
        if path:
            self.entry_template.delete(0, tk.END)
            self.entry_template.insert(0, path)

            # 如果输出目录还没填，默认用模板所在目录 / output_reports
            if not self.entry_output.get().strip():
                default_output = os.path.join(os.path.dirname(path), "output_reports")
                self.entry_output.insert(0, default_output)

    def browse_excel(self):
        path = filedialog.askopenfilename(
            title="选择 Excel 文件",
            filetypes=[("Excel 文件", "*.xlsx *.xls"), ("所有文件", "*.*")]
        )
        if path:
            self.entry_excel.delete(0, tk.END)
            self.entry_excel.insert(0, path)

    def browse_output(self):
        path = filedialog.askdirectory(title="选择输出文件夹")
        if path:
            self.entry_output.delete(0, tk.END)
            self.entry_output.insert(0, path)

    def log(self, msg):
        self.log_text.insert(tk.END, msg + "\n")
        self.log_text.see(tk.END)
        self.master.update_idletasks()

    def run(self):
        template_path = self.entry_template.get().strip()
        excel_path = self.entry_excel.get().strip()
        output_dir = self.entry_output.get().strip()

        if not template_path:
            messagebox.showwarning("提示", "请先选择模板文件。")
            return
        if not excel_path:
            messagebox.showwarning("提示", "请先选择 Excel 文件。")
            return
        if not output_dir:
            messagebox.showwarning("提示", "请先选择输出文件夹。")
            return

        # 清空日志
        self.log_text.delete(1.0, tk.END)

        try:
            self.log("开始生成报告...")
            generate_reports(template_path, excel_path, output_dir, log_func=self.log)
            messagebox.showinfo("完成", "所有报告生成完成！")
        except Exception as e:
            self.log("发生错误：")
            self.log(str(e))
            self.log(traceback.format_exc())
            messagebox.showerror("错误", f"生成过程中发生错误：{e}")


def main():
    root = tk.Tk()
    app = ReportFillerApp(root)
    root.mainloop()


if __name__ == "__main__":
    main()
